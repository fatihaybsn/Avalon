import sys
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen.canvas import Canvas
from PIL import Image, ImageDraw, ImageFont
root=Path(sys.argv[1])
d=Document();d.add_paragraph('Stage07 office fixture');d.save(root/'fixture.docx')
w=Workbook();w.active['A1']='=SUM(1,2)';w.save(root/'fixture.xlsx')
c=Canvas(str(root/'fixture.pdf'));c.drawString(40,800,'Stage07 PDF fixture');c.save()
Image.new('RGB',(16,16),'blue').save(root/'fixture.png')
ocr=Image.new('RGB',(1100,160),'white')
ImageDraw.Draw(ocr).text((25,35),'JARVIS OCR deneme 12345',font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',48),fill='black')
ocr.save(root/'ocr.png')
import zipfile
(root/'xml-secret.txt').write_text('JARVIS_SECRET_TEST',encoding='utf-8')
with zipfile.ZipFile(root/'fixture.pptx','w') as z:
    z.writestr('ppt/slides/slide1.xml', '<!DOCTYPE slide [<!ENTITY secret SYSTEM "'+(root/'xml-secret.txt').as_uri()+'">]><slide xmlns:a="urn:test"><a:t>Stage07 slide fixture</a:t><a:t>&secret;</a:t></slide>')
