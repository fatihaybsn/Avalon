"""Bounded local readers over installed libraries; no provider, macro or converter."""
import json
import sys
from pathlib import Path
sys.stdout.reconfigure(encoding='utf-8')

def read(p):
    path = Path(p['path'])
    if path.stat().st_size > 25 * 1024 * 1024:
        raise ValueError('Document exceeds local 25 MiB inspection budget')
    kind = p['kind']
    if kind == 'pdf':
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            parts = []
            truncated = False
            for page in pdf.pages[:20]:
                raw = page.extract_text() or ''
                truncated |= len(raw)>20000
                parts.append(raw[:20000])
                if sum(map(len, parts)) >= 20000:
                    break
            joined='\n'.join(parts)
            return dict(text=joined[:20000], pages=len(pdf.pages), partial=truncated or len(joined)>20000 or len(pdf.pages)>len(parts), ocrPerformed=False)
    if kind == 'office':
        # Preflight ZIP size before any XML reader (no macros/legacy executable formats).
        import zipfile
        with zipfile.ZipFile(path) as z:
            entries = z.infolist()
            if len(entries)>10000 or sum(e.file_size for e in entries)>50*1024*1024:
                raise ValueError('Office ZIP expansion budget exceeded')
            if any('vbaProject' in e.filename for e in entries):
                raise ValueError('Macro document refused')
            if path.suffix.lower() == '.docx':
                from docx import Document
                d=Document(path)
                return dict(text='\n'.join([a.text for a in d.paragraphs]+['\t'.join(c.text for c in r.cells) for t in d.tables for r in t.rows])[:20000], partial=True)
            if path.suffix.lower() == '.xlsx':
                import openpyxl
                book=openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
                try:
                    return dict(sheets=[dict(name=s.title, rows=[[str(v)[:1000] if v is not None else None for v in row] for row in s.iter_rows(min_row=1, max_row=min(s.max_row or 100,100), max_col=min(s.max_column or 20,20), values_only=True)]) for s in book.worksheets[:5]], partial=True, formulasExecuted=False)
                finally:
                    book.close()
            if path.suffix.lower() == '.pptx':
                from lxml.etree import fromstring, XMLParser
                parser=XMLParser(resolve_entities=False, no_network=True, load_dtd=False, huge_tree=False)
                slides=[e for e in entries if e.filename.startswith('ppt/slides/slide') and e.filename.endswith('.xml')][:30]
                return dict(text='\n'.join(' '.join(el.text or '' for el in fromstring(z.read(e),parser).iter() if isinstance(el.tag,str) and el.tag.endswith('}t')) for e in slides)[:20000], partial=True)
            raise ValueError('Use DOCX/XLSX/PPTX; legacy/macro formats need an explicitly selected installed adapter')
    if kind == 'image':
        from PIL import Image
        with Image.open(path) as im:
            return dict(format=im.format, size=list(im.size), mode=im.mode, frames=getattr(im,'n_frames',1), ocrPerformed=False)
    if kind == 'text':
        with path.open('r',encoding='utf-8') as f:
            data=f.read(20001)
        return dict(text=data[:20000],partial=len(data)>20000)
    raise ValueError('Unknown reader')

try:
    result=read(json.load(sys.stdin))
    if 'sheets' in result:
        while len(json.dumps(result,ensure_ascii=False))>40000:
            populated=[s for s in result['sheets'] if s['rows']]
            if not populated:
                break
            populated[-1]['rows'].pop()
    json.dump(dict(confirmed=True, trust='untrusted-document-data', authority='none', **result),sys.stdout,ensure_ascii=False)
except Exception as e:
    json.dump(dict(error=str(e)[:512],confirmed=False),sys.stdout)
    sys.exit(1)
